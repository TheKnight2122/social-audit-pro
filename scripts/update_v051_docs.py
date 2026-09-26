from docx import Document
from update_v050_docs import MANUAL, WEEKLY, add_table, add_bullets


TITLE = "Copias cifradas y recuperacion controlada"


def append_backup_section(document):
    if any(paragraph.text == TITLE for paragraph in document.paragraphs):
        return
    document.add_page_break()
    document.add_heading(TITLE, level=1)
    document.add_paragraph(
        "Avance 012 de la version v0.5.1. El sistema puede crear copias cifradas de SQLite y "
        "recuperarlas a un archivo nuevo sin sobrescribir la base activa. El recorrido se probo "
        "con datos desechables. Las copias periodicas siguen desactivadas hasta configurar una "
        "clave propia; no se cambiaron la base ni los secretos del usuario."
    )
    document.add_heading("Que incorpora este avance", level=2)
    add_table(document, ["Funcion", "Comportamiento"], [
        ["Copia online", "Incluye escrituras confirmadas en WAL usando la API de backup de SQLite."],
        ["Cifrado", "AES 256 GCM con clave independiente y metadatos autenticados."],
        ["Verificacion", "Descifrado, SHA 256, integridad, referencias y esquema compatible."],
        ["Restauracion", "Exige una ruta nueva e invalida sesiones, enlaces y codigos de recuperacion antiguos."],
        ["Programacion", "Intervalo opcional de 1 a 168 horas, sin solapamiento en el trabajador."],
    ], [1.55, 5.25])
    document.add_heading("Como utilizarlo", level=2)
    add_bullets(document, [
        "Configurar BACKUP_ENCRYPTION_KEY con 64 caracteres hexadecimales aleatorios y custodiar la clave fuera del servidor.",
        "Crear una copia: npm run db:backup. El resultado indica la carpeta generada.",
        "Comprobarla: npm run db:verify -- \"backups/CARPETA\".",
        "Recuperarla: npm run db:restore -- \"backups/CARPETA\" \"data/recuperada.sqlite\".",
        "Para copias diarias, establecer BACKUP_INTERVAL_HOURS=24 y reiniciar. Con 0 permanecen desactivadas.",
    ])
    document.add_paragraph(
        "CARPETA representa el nombre devuelto por la copia. Conservar tambien TOKEN_ENCRYPTION_KEY: "
        "es necesaria para OAuth y MFA despues de restaurar. Nunca guardar los valores reales en "
        "GitHub, capturas ni informes. La guia tecnica completa esta en docs/29-copias-y-restauracion.md."
    )
    document.add_page_break()
    document.add_heading("Recuperacion y evidencia de funcionamiento", level=1)
    document.add_paragraph(
        "Restaurar no cambia automaticamente la base en uso. El administrador del servidor detiene "
        "la aplicacion, conserva el archivo anterior, cambia DATABASE_PATH a la ruta recuperada y "
        "arranca una sola instancia. Debe revisar usuarios, roles, organizaciones y reportes antes "
        "de reabrir el servicio."
    )
    document.add_heading("Controles despues de recuperar", level=2)
    add_bullets(document, [
        "Las sesiones, enlaces temporales y estados OAuth previos dejan de ser validos.",
        "TOTP permanece activo; los codigos de recuperacion antiguos se eliminan y deben regenerarse desde Cuenta.",
        "Las sincronizaciones quedan pausadas hasta revisarlas y reactivarlas desde Integraciones.",
        "Las contrasenas y los permisos vuelven al estado de la copia: revisar cambios o revocaciones posteriores.",
    ])
    document.add_heading("Pruebas realizadas", level=2)
    add_table(document, ["Verificacion", "Resultado"], [
        ["Bateria completa", "34 pruebas aprobadas, incluidas 9 de respaldo y recuperacion."],
        ["Recorrido CLI", "Crear, verificar y restaurar funcionan con una base desechable."],
        ["Casos negativos", "Clave incorrecta, archivo alterado, referencias invalidas y destino existente rechazados."],
        ["Programacion", "Repeticion por intervalo, no solapamiento y espera al cerrar comprobados."],
        ["Base original", "No se modifica durante la restauracion; la prueba verifica que conserva sus sesiones y tareas."],
    ], [1.8, 5.0])
    document.add_heading("Lo que aun falta", level=2)
    document.add_paragraph(
        "Las copias locales no protegen frente a perder el equipo completo. Faltan almacenamiento "
        "externo, retencion, alertas y simulacros productivos. Un cierre forzado puede dejar "
        "instantaneas temporales: proteger permisos y cifrado del disco. No hay borrado automatico "
        "de copias. PostgreSQL, cola compartida, despliegue productivo y validacion con cuentas "
        "sociales reales siguen pendientes. No se contrato ningun servicio."
    )


for path in [MANUAL, WEEKLY]:
    document = Document(path)
    document.core_properties.author = "Social Audit Pro Team"
    document.core_properties.last_modified_by = "Social Audit Pro Team"
    # Preserve the previous version's evidence and append this advance chronologically.
    if path == MANUAL:
        for table in document.tables:
            if table.rows and table.rows[0].cells[0].text == "Dato":
                for row in table.rows:
                    if row.cells[0].text == "Version documentada":
                        row.cells[1].text = "v0.5.1"
                    if row.cells[0].text == "Referencia de avance":
                        row.cells[1].text = "avance-012"
                break
    else:
        for paragraph in document.paragraphs:
            if paragraph.text.startswith("Version documentada"):
                paragraph.text = paragraph.text.replace("v0.5.0", "v0.5.1")
                break
    append_backup_section(document)
    document.save(path)
    print(path)
