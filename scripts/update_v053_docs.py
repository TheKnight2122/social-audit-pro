from docx import Document
from update_v050_docs import MANUAL, WEEKLY, add_table


TITLE = "Sincronizacion segura y pruebas de carga local"


def append_validation(document):
    if any(paragraph.text == TITLE for paragraph in document.paragraphs):
        return
    document.add_page_break()
    document.add_heading(TITLE, level=1)
    document.add_paragraph(
        "Avance 014 de la version v0.5.3 del 26 de septiembre de 2026. Se reforzo el control "
        "de sincronizaciones simultaneas y la separacion de empresas. Las 58 pruebas automatizadas "
        "pasaron. Una carga local de 500 solicitudes con 10 simultaneas termino sin errores "
        "y guardo correctamente 50 reportes en una base desechable."
    )
    add_table(document, ["Situacion", "Comportamiento comprobado"], [
        ["Dos sincronizaciones", "La manual y la automatica comparten bloqueo; la segunda manual recibe 409 mientras este ocupado."],
        ["Consulta lenta", "El bloqueo se renueva. Al superar 120 segundos, el resultado tardio no se guarda."],
        ["Cambio de acceso", "Revocar la conexion, desactivar la empresa o retirar permisos impide guardar resultados en curso."],
        ["Cuenta de otra empresa", "Los accesos cruzados a reportes, historicos, publicaciones, usuarios y actividad son rechazados o filtrados."],
        ["Retorno de OAuth", "El estado se consume una sola vez; se revalidan permisos y se evita sobrescribir una sincronizacion activa."],
    ], [1.6, 5.2])
    document.add_heading("Como verificarlo", level=2)
    document.add_paragraph(
        "npm test ejecuta las 58 pruebas. npm run test:load -- 500 10 inicia un servidor "
        "temporal propio, hace consultas autenticadas y crea reportes ficticios. No acepta una "
        "pagina externa como destino, no carga los secretos de .env ni usa la base habitual. "
        "Al terminar cierra el servidor y elimina su carpeta temporal."
    )
    document.add_paragraph(
        "GitHub Actions incorpora una carga corta antes de publicar. La evidencia local esta "
        "en docs/evidencias/carga-v053.json. La guia tecnica y los limites se encuentran en "
        "docs/31-validacion-concurrencia-y-carga.md."
    )
    document.add_heading("Que no demuestra esta prueba", level=2)
    document.add_paragraph(
        "La carga usa pocos datos y sesiones ya creadas; no mide login, correo ni APIs sociales. "
        "No determina cuantos clientes soportara un servidor real ni reemplaza una auditoria "
        "externa. Si un SDK ignora cancelacion puede terminar una consulta externa, aunque su "
        "resultado ya no se escriba en la base. SQLite sigue siendo de una instancia."
    )
    document.add_heading("Trabajo restante", level=2)
    document.add_paragraph(
        "La solicitud general aun no esta completada. Faltan credenciales y aprobaciones de redes, "
        "validacion real, SMTP y despliegue productivo. Tambien siguen pendientes PostgreSQL, cola "
        "compartida, balanceador, copias externas, monitoreo centralizado y pruebas sostenidas. "
        "No se contrataron servicios ni se cambiaron las pantallas; la pagina publica sigue "
        "siendo una demostracion estatica."
    )


def main():
    for path in [MANUAL, WEEKLY]:
        document = Document(path)
        document.core_properties.author = "Social Audit Pro Team"
        document.core_properties.last_modified_by = "Social Audit Pro Team"
        if path == MANUAL:
            for table in document.tables:
                if table.rows and table.rows[0].cells[0].text == "Dato":
                    for row in table.rows:
                        if row.cells[0].text == "Version documentada":
                            row.cells[1].text = "v0.5.3"
                        if row.cells[0].text == "Referencia de avance":
                            row.cells[1].text = "avance-014"
                    break
        else:
            for paragraph in document.paragraphs:
                if paragraph.text.startswith("Version documentada"):
                    paragraph.text = "Version documentada  v0.5.3  |  Fecha de corte  26 de septiembre de 2026"
                    break
        append_validation(document)
        document.save(path)
        print(path)


if __name__ == "__main__":
    main()
