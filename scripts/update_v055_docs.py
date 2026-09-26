from pathlib import Path
from docx import Document
from docx.shared import Inches
from update_v050_docs import MANUAL, WEEKLY

TITLE = "Microsoft 365 y pruebas sin contratar alojamiento"


def main():
    for path in [MANUAL, WEEKLY]:
        document = Document(path)
        if any(p.text == TITLE for p in document.paragraphs):
            continue
        document.core_properties.author = "Social Audit Pro Team"
        document.core_properties.last_modified_by = "Social Audit Pro Team"
        if path == MANUAL:
            for table in document.tables:
                if table.rows and table.rows[0].cells[0].text == "Dato":
                    for row in table.rows:
                        if row.cells[0].text == "Version documentada":
                            row.cells[1].text = "v0.5.5"
                        if row.cells[0].text == "Referencia de avance":
                            row.cells[1].text = "avance-016"
                    break
        else:
            for paragraph in document.paragraphs:
                if paragraph.text.startswith("Version documentada"):
                    paragraph.text = "Version documentada  v0.5.5  |  Fecha de corte  26 de septiembre de 2026"
                    break
        document.add_page_break()
        document.add_heading(TITLE, level=1)
        document.add_paragraph(
            "Avance 016 del 26 de septiembre de 2026. Se confirmo Microsoft 365, administrado "
            "por el jefe, y un piloto gratuito con pausas por inactividad aceptadas. Las empresas "
            "deben ver solo sus cuentas designadas y creceran gradualmente, sin un volumen inicial "
            "definido. Se pide iniciar pruebas cuanto antes y mantener desactivadas las APIs de pago."
        )
        document.add_heading("Lo implementado", level=2)
        document.add_paragraph(
            "Se preparo el envio de verificaciones y recuperaciones mediante Microsoft Graph HTTPS. "
            "Usa credenciales de aplicacion y un buzon remitente fijo, no la contrasena del jefe. "
            "El token temporal se conserva solo en memoria. Las solicitudes tienen limite de "
            "tiempo y no siguen redirecciones. Los errores persistidos no guardan mensajes "
            "originales que puedan incluir claves o contenido privado. SMTP sigue disponible."
        )
        document.add_paragraph(
            "X queda bloqueado por defecto aun si existen credenciales guardadas. El bloqueo "
            "abarca OAuth, sincronizacion manual, dashboard y trabajador programado. La interfaz "
            "indica Desactivada por presupuesto. No se contrataron servicios ni se activaron "
            "cuentas externas. Las 76 pruebas automatizadas pasaron con respuestas simuladas."
        )
        document.add_heading("Lo que aun no esta activo", level=2)
        document.add_paragraph(
            "No se ha verificado un correo real. El jefe debe autorizar la aplicacion y limitar "
            "su permiso al buzon designado. La aceptacion HTTP 202 de Graph no garantiza entrega "
            "al destinatario. El sistema no reintenta automaticamente el envio para evitar duplicados. "
            "Faltan tambien PostgreSQL, backend online y validacion con redes reales. "
            "GitHub Pages sigue siendo una demo, no el sistema completo alojado."
        )
        document.add_heading("Asignacion por empresa", level=2)
        document.add_paragraph(
            "El administrador crea o selecciona la organizacion, conecta sus cuentas y registra "
            "los usuarios Cliente dentro de esa empresa. No se deben mezclar empresas distintas "
            "en una misma organizacion. La asignacion actual no incluye compartir una cuenta "
            "entre empresas ni restringir usuarios individuales de una misma empresa."
        )
        document.add_page_break()
        document.add_heading("Pasos para iniciar el piloto", level=1)
        document.add_paragraph(
            "El jefe de Microsoft 365 designa el buzon y autoriza la aplicacion con envio limitado "
            "mediante Exchange RBAC. El administrador prepara las cuentas gratuitas del alojamiento "
            "y de la base de datos. Los identificadores y secretos se configuran de forma privada "
            "en el servidor; nunca se publican en GitHub ni se envian contrasenas por el chat."
        )
        document.add_paragraph(
            "Antes del piloto online se debe portar SQLite a PostgreSQL y validar permisos, "
            "correo, login, aislamiento y consultas reales con pocos usuarios. Las pausas y cuotas "
            "del plan gratuito no equivalen a continuidad garantizada. Toda API debe revisarse "
            "antes de activarla; el sistema no detecta cambios de tarifas externos."
        )
        document.add_picture(str(Path(__file__).resolve().parents[1] / "docs/assets/piloto/x-bloqueado-desktop.png"), width=Inches(4.8))
        document.add_paragraph(
            "Captura de verificacion local con datos desechables. La API de X permanece bloqueada "
            "por presupuesto. Se reviso la misma vista en escritorio y movil."
        )
        document.add_paragraph(
            "Guia operativa y fuentes oficiales: docs/33-piloto-microsoft365-y-presupuesto.md. "
            "La guia especifica las variables del servidor y las comprobaciones de alcance "
            "que debe realizar el administrador de Microsoft 365."
        )
        document.save(path)
        print(path)


if __name__ == "__main__":
    main()
