from docx import Document
from update_v050_docs import MANUAL, WEEKLY

TITLE = "Recuperacion de cuenta preparada para PostgreSQL"


def main():
    for path in [MANUAL, WEEKLY]:
        document = Document(path)
        if any(p.text == TITLE for p in document.paragraphs):
            continue
        document.core_properties.author = "Social Audit Pro Team"
        document.core_properties.last_modified_by = "Social Audit Pro Team"
        if path == MANUAL:
            for table in document.tables:
                if table.rows and table.rows[0].cells[0].text == "Dato":
                    for row in table.rows:
                        if row.cells[0].text == "Version documentada":
                            row.cells[1].text = "v0.5.8"
                        if row.cells[0].text == "Referencia de avance":
                            row.cells[1].text = "avance-019"
                    break
        else:
            for paragraph in document.paragraphs:
                if paragraph.text.startswith("Version documentada"):
                    paragraph.text = "Version documentada  v0.5.8  |  Fecha de corte  26 de septiembre de 2026"
                    break
        document.add_page_break()
        document.add_heading(TITLE, level=1)
        document.add_paragraph(
            "Avance 019 del 26 de septiembre de 2026. La recuperacion de contrasena y la "
            "verificacion de correo ya se prueban con los mismos endpoints en SQLite y PostgreSQL. "
            "El servidor local completo sigue usando SQLite; no se ha desplegado un backend nuevo "
            "ni se ha activado el correo empresarial."
        )
        document.add_heading("Que cambia para la seguridad", level=2)
        document.add_paragraph(
            "Cada nuevo enlace sustituye al anterior de su mismo tipo. Solo se puede completar "
            "una vez, incluso con solicitudes simultaneas. El sistema rechaza enlaces vencidos y "
            "cuentas desactivadas. Al restablecer la contrasena se cierran las sesiones y se "
            "invalidan los desafios de acceso pendientes, dentro de una sola transaccion."
        )
        document.add_paragraph(
            "El segundo factor ya activo se conserva; una configuracion de 2FA que todavia no "
            "se habia confirmado se elimina. El registro de correos conserva destinatario, asunto "
            "y estado, sin guardar el cuerpo ni enlaces secretos. Si falla la base despues de que "
            "el proveedor acepte un correo, no se reenvia automaticamente ni se atribuye el fallo "
            "al proveedor. La aceptacion tampoco garantiza llegada al buzon."
        )
        document.add_heading("Pruebas realizadas", level=2)
        document.add_paragraph(
            "Pasaron 105 pruebas automatizadas, la construccion publica y 200 solicitudes de "
            "carga local sin errores. Los contratos comprueban recuperacion, verificacion, "
            "caducidad, aislamiento, rollback y privacidad. La prueba PostgreSQL del workflow "
            "incluye solicitudes concurrentes desde dos grupos de conexiones. Todos los correos "
            "de prueba son capturados o simulados; no se envio correo real."
        )
        document.add_heading("Lo que queda pendiente", level=2)
        document.add_paragraph(
            "Faltan registro, login y gestion 2FA completos sobre PostgreSQL, ademas de los "
            "modulos de datos y trabajadores. Esas funciones siguen disponibles localmente con "
            "SQLite. La publica permanece como demo estatica. Siguen pendientes autorizaciones "
            "sociales, Microsoft 365, alojamiento y pruebas reales. No se contrataron servicios. "
            "Detalle tecnico: docs/36-recuperacion-postgresql.md."
        )
        document.save(path)
        print(path)


if __name__ == "__main__":
    main()
