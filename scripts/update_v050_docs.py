from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
MANUAL = ROOT / "docs" / "manuales" / "Manual-de-avance-Social-Audit-Pro.docx"
WEEKLY = ROOT / "docs" / "semanales" / "Informe-semanal-2026-09-21-al-27.docx"
ASSETS = ROOT / "docs" / "assets" / "manual"

BLACK = "000000"
NAVY = "17324D"
TEAL = "2F7D78"
PALE = "EDF5F4"
PALE_BLUE = "EEF3F8"
GRAY = "5C6670"
BORDER = "D9D9D9"


def set_cell_fill(cell, color):
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), color)


def set_cell_margins(cell, top=120, start=140, bottom=120, end=140):
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn("w:" + side))
        if node is None:
            node = OxmlElement("w:" + side)
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    properties = table._tbl.tblPr
    borders = properties.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn("w:" + edge))
        if node is None:
            node = OxmlElement("w:" + edge)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:color"), BORDER)


def keep_row_together(row):
    properties = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    properties.append(cant_split)


def shade_table(table, widths=None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    for row_index, row in enumerate(table.rows):
        keep_row_together(row)
        for column_index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if widths:
                cell.width = Inches(widths[column_index])
            set_cell_fill(cell, NAVY if row_index == 0 else (PALE_BLUE if row_index % 2 == 0 else "FFFFFF"))
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(9.2)
                    if row_index == 0:
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)


def add_table(document, headers, rows, widths):
    table = document.add_table(rows=1, cols=len(headers))
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for index, value in enumerate(headers):
        table.rows[0].cells[index].text = value
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = value
    shade_table(table, widths)
    return table


def configure_styles(document):
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(35, 42, 48)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.12
    for name, size in (("Title", 28), ("Heading 1", 18), ("Heading 2", 13)):
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(14 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(7)
        paragraph_properties = style.element.get_or_add_pPr()
        border = paragraph_properties.find(qn("w:pBdr"))
        if border is not None:
            paragraph_properties.remove(border)
    styles["List Bullet"].font.name = "Arial"
    styles["List Bullet"].font.size = Pt(10.5)


def configure_page(document):
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.text = "Social Audit Pro  |  Documentacion del proyecto"
    for run in footer.runs:
        run.font.name = "Arial"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(92, 102, 112)


def add_label(document, text):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(text.upper())
    run.font.name = "Arial"
    run.font.size = Pt(9)
    run.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)
    return paragraph


def add_bullets(document, items):
    for item in items:
        paragraph = document.add_paragraph(item, style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(4)


def add_figure(document, image_path, caption, width=6.65):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(image_path), width=Inches(width))
    caption_paragraph = document.add_paragraph(caption)
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.paragraph_format.space_before = Pt(3)
    caption_paragraph.paragraph_format.space_after = Pt(10)
    for run in caption_paragraph.runs:
        run.italic = True
        run.font.name = "Arial"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(92, 102, 112)


def update_manual():
    document = Document(MANUAL)
    configure_styles(document)
    configure_page(document)
    document.core_properties.author = "Social Audit Pro Team"
    document.core_properties.last_modified_by = "Social Audit Pro Team"
    replacements = {
        "v0.4.0": "v0.5.0",
        "20 de septiembre de 2026": "25 de septiembre de 2026",
        "Commit de referencia": "Referencia de avance",
        "ca1d24b": "avance-011",
        "MVP avanzado con primer OAuth oficial": "MVP con organizaciones seguridad y sincronizacion",
        "17 aprobadas": "25 aprobadas",
        "Aislamiento OAuth por usuario": "Aislamiento por organizacion",
        "Aislamiento entre usuarios": "Aislamiento entre organizaciones",
        "Sincronizacion manual y actualizacion oficial de tokens.": "Sincronizacion manual y programada con renovacion de tokens cuando el proveedor la admite.",
        "Sustitucion automatica de datos demo por datos oficiales para el usuario conectado.": "Sustitucion de datos demo por datos oficiales de la organizacion activa.",
        "Permite conectar YouTube mediante OAuth oficial y sincronizar datos disponibles.": "Incluye conectores para seis redes y sincronizacion programada, sujetos a credenciales y aprobacion.",
        "Conclusion principal: el MVP ya es funcional como sistema local y YouTube es la primera red con un conector OAuth oficial completo. Las demas redes permanecen preparadas visualmente, pero no se presentan como conectadas hasta disponer de sus credenciales y permisos oficiales.": "Conclusion principal: el MVP local incluye organizaciones, MFA, correo, recuperacion y sincronizacion automatica. Los seis conectores estan implementados; la activacion y validacion real requieren credenciales y permisos oficiales.",
        "Las capturas siguientes corresponden a la version actual. Las cifras visibles en estas imagenes son demostrativas porque no se incluyeron credenciales personales en el repositorio.": "Estas capturas conservan el recorrido visual de la linea base v0.4.0. El suplemento final explica las funciones de v0.5.0. Las cifras son demostrativas y no incluyen datos privados.",
        "La vista muestra el estado real de cada proveedor. YouTube dispone del flujo oficial; los demas proveedores permanecen sin configurar hasta contar con sus aplicaciones y permisos aprobados.": "La vista muestra el estado de cada proveedor. Los seis conectores requieren aplicaciones y permisos aprobados para activarse. Los cinco adaptadores nuevos se validaron con respuestas simuladas.",
        "YouTube es la primera integracion implementada de extremo a extremo para validar el patron que utilizaran los proveedores restantes.": "YouTube fue la primera integracion y establecio el patron compartido por los seis adaptadores actuales.",
        "El avance de YouTube fue validado con pruebas automaticas y revision visual. Las pruebas del proveedor usan una simulacion controlada para comprobar la aplicacion sin enviar credenciales reales a Google.": "Las pruebas automaticas verifican analitica, API, seguridad, organizaciones, OAuth y sincronizacion. Los proveedores usan respuestas simuladas; falta validar cuentas reales con credenciales propias.",
        "Commit actual: ca1d24b feat integrar OAuth oficial de YouTube.": "Referencia actual: docs/avances/avance-011.md y CHANGELOG v0.5.0.",
        "La siguiente etapa no debe reconstruir el dashboard. Debe reutilizar el patron validado con YouTube y avanzar por proveedor conforme se obtengan credenciales y permisos oficiales.": "La siguiente etapa requiere activar y validar los proveedores, configurar SMTP y elegir infraestructura productiva con base compartida.",
        "No reemplazar SQLite durante el MVP sin una necesidad demostrada.": "Migrar SQLite a una base compartida antes de ejecutar multiples servidores.",
        "Esta actualizacion reemplaza el estado anterior cuando exista una diferencia. El sistema local ahora incorpora organizaciones, seguridad de cuenta, sincronizacion programada y adaptadores para las seis redes. La demostracion de GitHub Pages continua siendo estatica y no procesa datos privados.": "El sistema local incorpora organizaciones, seguridad de cuenta, sincronizacion programada y adaptadores para las seis redes. Este suplemento detalla el avance v0.5.0. GitHub Pages sigue siendo una demostracion estatica.",
    }
    paragraphs = list(document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    for paragraph in paragraphs:
        # La tabla historica conserva las versiones previas.
        if paragraph.text in replacements and not (paragraph.text == "v0.4.0" and paragraph._p.getparent().getparent().getparent() is document.tables[8]._tbl):
            paragraph.text = replacements[paragraph.text]
    for row in document.tables[1].rows:
        if row.cells[0].text == "Otras redes":
            row.cells[1].text = "Codigo implementado"
    for row in document.tables[6].rows:
        if row.cells[0].text in ("Recuperacion de contrasena y MFA", "Organizaciones y clientes"):
            row.cells[1].text = "Implementado"
    for row in document.tables[7].rows:
        if row.cells[0].text == "Auditoria npm":
            row.cells[1].text = "Revision adicional requerida"
    pending = [
        ("1", "Validar YouTube con una cuenta real", "Credenciales Google"),
        ("2", "Activar Facebook e Instagram", "App Meta aprobada"),
        ("3", "Activar TikTok LinkedIn y X", "Acceso oficial y cuotas"),
        ("4", "Validar tareas con proveedores reales", "Credenciales y limites"),
        ("5", "Completar desconexion y revocacion", "Endpoints de cada red"),
        ("6", "Migrar a PostgreSQL y cola compartida", "Infraestructura elegida"),
        ("7", "Activar SMTP y probar entrega", "Servicio de correo"),
        ("8", "Desplegar backend en HTTPS", "Host dominio y secretos"),
    ]
    for row, values in zip(document.tables[10].rows[1:], pending):
        for cell, value in zip(row.cells, values):
            cell.text = value
    if any(paragraph.text == "Actualizacion v0.5.0" for paragraph in document.paragraphs):
        document.save(MANUAL)
        return
    document.add_page_break()

    add_label(document, "Actualizacion del estado")
    document.add_heading("Actualizacion v0.5.0", level=1)
    document.add_paragraph(
        "Esta actualizacion reemplaza el estado anterior cuando exista una diferencia. "
        "El sistema local ahora incorpora organizaciones, seguridad de cuenta, sincronizacion "
        "programada y adaptadores para las seis redes. La demostracion de GitHub Pages continua "
        "siendo estatica y no procesa datos privados."
    )

    document.add_heading("Capacidades incorporadas", level=2)
    add_table(
        document,
        ["Area", "Estado en v0.5.0", "Evidencia"],
        [
            ("Organizaciones", "Membresias, seleccion de contexto y filtros por organizacion", "Migracion 003 y rutas de organizaciones"),
            ("Seguridad", "Verificacion, recuperacion y MFA TOTP", "Migracion 005 y pruebas de cuenta"),
            ("Sincronizacion", "Intervalos configurables y trabajador con arrendamiento", "sync-worker.js y prueba de ejecucion unica"),
            ("Conectores", "YouTube, Instagram, Facebook, TikTok, LinkedIn y X", "Registro de proveedores y pruebas de normalizacion"),
            ("Operacion", "Docker, salud y apagado ordenado", "Dockerfile, Compose y rutas de health"),
        ],
        [1.35, 3.15, 2.15],
    )

    document.add_heading("Estado de los conectores", level=2)
    document.add_paragraph(
        "El codigo de integracion no equivale a una conexion activa. Cada plataforma exige una "
        "aplicacion registrada, permisos aprobados, credenciales privadas y validacion con una "
        "cuenta real."
    )
    add_table(
        document,
        ["Plataforma", "Implementacion", "Activacion externa"],
        [
            ("YouTube", "OAuth, renovacion y normalizacion", "Credenciales de Google"),
            ("Instagram", "OAuth y perfil profesional", "App, permisos y revision de Meta"),
            ("Facebook", "OAuth, paginas y publicaciones", "App, permisos y revision de Meta"),
            ("TikTok", "OAuth PKCE, perfil y videos", "App y productos aprobados"),
            ("LinkedIn", "OAuth, pagina y estadisticas", "Community Management API"),
            ("X", "OAuth PKCE, perfil y publicaciones", "Proyecto, plan y cuotas"),
        ],
        [1.25, 2.65, 2.75],
    )

    document.add_page_break()
    document.add_heading("Flujo de seguridad de cuenta", level=1)
    add_bullets(document, [
        "El correo debe confirmarse antes de iniciar una sesion nueva.",
        "Los enlaces de verificacion y recuperacion expiran, se consumen una vez y se guardan como hash.",
        "El cambio de contrasena elimina sesiones anteriores.",
        "El segundo factor usa TOTP, secreto cifrado y diez codigos de recuperacion con hash.",
        "Los tokens OAuth y verificadores PKCE nunca se envian al frontend.",
    ])

    document.add_heading("Version local y version publica", level=2)
    add_table(
        document,
        ["Caracteristica", "Local", "Publica"],
        [
            ("Backend y API", "Si", "No"),
            ("Base de datos", "SQLite persistente", "No almacena datos"),
            ("Usuarios y MFA", "Funcionales", "Solo presentacion visual"),
            ("OAuth y sincronizacion", "Disponibles con credenciales", "Deshabilitados"),
            ("Actualizacion", "Al ejecutar el codigo local", "Tras enviar a main y completar GitHub Actions"),
        ],
        [2.4, 2.1, 2.15],
    )

    document.add_heading("Produccion y alta disponibilidad", level=2)
    document.add_paragraph(
        "La aplicacion esta preparada para empaquetarse como una instancia, pero no se considera "
        "desplegada en produccion. SQLite impide operar multiples servidores de forma segura. "
        "La siguiente fase debe incorporar PostgreSQL administrado, una cola compartida, balanceo, "
        "copias, observabilidad y pruebas de recuperacion."
    )

    document.add_heading("Validacion de esta version", level=2)
    add_bullets(document, [
        "25 pruebas automatizadas aprobadas.",
        "Compilacion estatica completada.",
        "Cuenta y MFA revisados en escritorio y movil.",
        "No se incorporaron credenciales reales ni datos privados.",
        "Docker no estaba instalado en el equipo, por lo que el contenedor no se ejecuto localmente.",
    ])

    output = MANUAL.with_suffix(".tmp.docx")
    document.save(output)
    output.replace(MANUAL)


def create_weekly():
    document = Document()
    configure_styles(document)
    configure_page(document)
    document.core_properties.title = "Informe semanal Social Audit Pro 21 al 27 de septiembre de 2026"
    document.core_properties.subject = "Estado verificable del desarrollo y pendientes"
    document.core_properties.author = "Social Audit Pro Team"
    document.core_properties.last_modified_by = "Social Audit Pro Team"

    add_label(document, "Social Audit Pro")
    title = document.add_paragraph("Informe semanal del proyecto", style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    subtitle = document.add_paragraph("Semana del 21 al 27 de septiembre de 2026")
    subtitle.paragraph_format.space_after = Pt(16)
    for run in subtitle.runs:
        run.font.name = "Arial"
        run.font.size = Pt(14)
        run.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)
    document.add_paragraph(
        "Este informe presenta el avance verificable de la plataforma, explica que funciones "
        "quedaron implementadas y separa con precision el trabajo que aun depende de credenciales, "
        "aprobaciones o infraestructura externa."
    )
    conclusion = document.add_paragraph()
    lead = conclusion.add_run("Conclusion principal. ")
    lead.bold = True
    conclusion.add_run(
        "La version local ya cuenta con organizaciones, recuperacion de contrasena, verificacion "
        "de correo, MFA, sincronizacion automatica y conectores para las seis redes. Todavia no es "
        "un servicio productivo de alta disponibilidad."
    )
    document.add_paragraph("Version documentada  v0.5.0  |  Fecha de corte  25 de septiembre de 2026")

    document.add_page_break()
    document.add_heading("Solicitudes atendidas", level=1)
    document.add_paragraph(
        "La solicitud de la semana concentro seguridad, multiempresa, integraciones oficiales, "
        "automatizacion y preparacion de produccion."
    )
    add_table(
        document,
        ["Solicitud", "Resultado", "Estado"],
        [
            ("Instagram Facebook TikTok LinkedIn y X", "Adaptadores OAuth y normalizacion implementados", "Falta activacion externa"),
            ("Sincronizacion automatica", "Programaciones, intervalos y trabajador con arrendamiento", "Implementado"),
            ("Backend base y autenticacion en produccion", "Contenedor, validaciones y health checks", "Falta proveedor productivo"),
            ("Recuperacion verificacion y 2FA", "Flujos completos con correo, TOTP y recuperacion", "Implementado"),
            ("Aislamiento entre empresas", "Organizaciones, membresias y filtros por contexto", "Implementado y sujeto a revision"),
            ("Multiples servidores", "Bases tecnicas documentadas", "Falta PostgreSQL y cola"),
        ],
        [2.15, 3.25, 1.25],
    )

    document.add_heading("Resultado funcional", level=1)
    add_bullets(document, [
        "Un usuario puede pertenecer a organizaciones y cambiar el contexto activo.",
        "Los datos de cuentas, analitica, reportes, actividad e integraciones se filtran por organizacion.",
        "El correo se verifica antes de un nuevo login y la contrasena puede restablecerse con enlace temporal.",
        "MFA permite alta por QR, codigos TOTP y codigos de recuperacion.",
        "Las conexiones pueden sincronizarse manualmente o con intervalos programados.",
        "Los seis proveedores comparten una ruta OAuth y un contrato normalizado de datos.",
    ])

    add_figure(
        document,
        ASSETS / "integraciones-desktop.png",
        "Figura 1  Vista de Integraciones utilizada para administrar proveedores",
    )

    document.add_page_break()
    document.add_heading("Arquitectura alcanzada", level=1)
    document.add_paragraph(
        "El frontend consume una API Express. La API aplica identidad, permisos y contexto de "
        "organizacion antes de consultar SQLite. Los adaptadores externos normalizan respuestas y "
        "el trabajador procesa las tareas programadas."
    )
    add_figure(
        document,
        ASSETS / "arquitectura.png",
        "Figura 2  Capas principales de la aplicacion local",
        width=6.35,
    )
    add_table(
        document,
        ["Componente", "Responsabilidad"],
        [
            ("Frontend", "Navegacion, formularios, filtros y presentacion de resultados"),
            ("API", "Autorizacion, validacion, orquestacion y respuestas REST"),
            ("Identidad", "Sesiones, correo, recuperacion, MFA, CSRF y roles"),
            ("Integraciones", "OAuth, renovacion, consulta y normalizacion por red"),
            ("Trabajador", "Reclamo y ejecucion de sincronizaciones vencidas"),
            ("Persistencia", "Organizaciones, historicos, tokens cifrados, reportes y actividad"),
        ],
        [1.75, 4.9],
    )

    document.add_page_break()
    document.add_heading("Seguridad y aislamiento", level=1)
    add_table(
        document,
        ["Control", "Implementacion", "Riesgo restante"],
        [
            ("Correo", "Token con hash, expiracion y uso unico", "Configurar SMTP y dominio"),
            ("Contrasenas", "scrypt y revocacion de sesiones al restablecer", "Agregar monitoreo de abuso"),
            ("MFA", "TOTP cifrado y codigos de recuperacion con hash", "Probar recuperacion operativa"),
            ("OAuth", "Estado de un uso, cifrado y PKCE cuando aplica", "Validar revocacion por proveedor"),
            ("Organizaciones", "Contexto en sesion y filtros de recursos", "Revision externa de autorizacion"),
            ("Produccion", "HTTPS y secreto validados al iniciar", "Gestor externo de secretos"),
        ],
        [1.35, 3.05, 2.25],
    )

    document.add_heading("Pruebas ejecutadas", level=2)
    add_bullets(document, [
        "25 pruebas automatizadas ejecutadas y aprobadas.",
        "OAuth generico PKCE comprobado desde el inicio hasta el token cifrado.",
        "Normalizacion simulada comprobada para TikTok, X, Facebook, Instagram y LinkedIn.",
        "Dos trabajadores comprobados contra una misma tarea para evitar ejecucion duplicada.",
        "Aislamiento entre dos organizaciones comprobado en analitica y reportes.",
        "Compilacion de la demostracion publica completada.",
    ])

    document.add_heading("Datos y privacidad", level=2)
    document.add_paragraph(
        "No se agregaron secretos reales, nombres personales ni sesiones al repositorio. La "
        "demostracion de GitHub Pages conserva datos de muestra y no puede acceder a la base local."
    )

    document.add_page_break()
    document.add_heading("Conectores y activacion externa", level=1)
    add_table(
        document,
        ["Red", "Listo en codigo", "Necesario para activarla"],
        [
            ("YouTube", "OAuth, consulta, renovacion y persistencia", "Cliente OAuth y APIs de Google"),
            ("Instagram", "OAuth y perfil profesional", "App de Meta y permisos aprobados"),
            ("Facebook", "Paginas, publicaciones y token de pagina", "App de Meta y permisos aprobados"),
            ("TikTok", "PKCE, perfil, videos y renovacion", "Aplicacion y productos aprobados"),
            ("LinkedIn", "Pagina, seguidores y estadisticas", "Community Management API"),
            ("X", "PKCE, perfil, publicaciones y renovacion", "Proyecto, plan y cuotas"),
        ],
        [1.15, 3.05, 2.45],
    )
    document.add_paragraph(
        "La implementacion se considera preparada para activacion, no conectada. Los proveedores "
        "pueden modificar scopes, versiones y limites; por eso cada alta requiere una revision de "
        "la documentacion oficial y una prueba con cuenta autorizada."
    )

    document.add_heading("Produccion y alta disponibilidad", level=1)
    document.add_paragraph(
        "El Dockerfile y el Compose permiten preparar una instancia. Para multiples servidores, "
        "SQLite debe sustituirse por PostgreSQL o equivalente y las tareas deben pasar a una cola "
        "compartida. Tambien faltan balanceo, observabilidad, copias y pruebas de restauracion."
    )
    add_table(
        document,
        ["Paso siguiente", "Motivo", "Prioridad"],
        [
            ("Elegir proveedor productivo", "Define dominio, red, base, correo y secretos", "Alta"),
            ("Migrar a PostgreSQL", "Permite persistencia compartida y replicas", "Alta"),
            ("Validar un conector adicional", "Reduce riesgo antes de activar las cinco redes", "Alta"),
            ("Incorporar cola y observabilidad", "Controla trabajos, reintentos y fallos", "Media"),
            ("Pruebas de carga y seguridad", "Confirma capacidad y controles antes de clientes", "Alta"),
        ],
        [2.35, 3.25, 1.05],
    )

    document.add_heading("Fuentes de comprobacion", level=2)
    add_bullets(document, [
        "Migraciones 003 a 006.",
        "Rutas de organizaciones, autenticacion e integraciones.",
        "Adaptadores de proveedores y trabajador de sincronizacion.",
        "Pruebas de API, cuenta, proveedores y trabajador.",
        "README, CHANGELOG y documentos 25 a 28.",
    ])

    WEEKLY.parent.mkdir(parents=True, exist_ok=True)
    document.save(WEEKLY)


if __name__ == "__main__":
    update_manual()
    create_weekly()
    print(MANUAL)
    print(WEEKLY)
