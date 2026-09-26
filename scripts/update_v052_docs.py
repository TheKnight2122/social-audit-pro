from docx import Document
from update_v050_docs import MANUAL, WEEKLY, add_table, add_bullets


TITLE = "Monitoreo operativo del servidor"


def append_monitoring(document):
    if any(paragraph.text == TITLE for paragraph in document.paragraphs):
        return
    document.add_page_break()
    document.add_heading(TITLE, level=1)
    document.add_paragraph(
        "Avance 013 de la version v0.5.2, registrado el 26 de septiembre de 2026. "
        "El backend permite detectar errores y lentitud mediante registros de solicitudes y "
        "estadisticas del proceso. No cambia las pantallas ni contrata servicios. "
        "Las 42 pruebas automatizadas pasaron, incluidas ocho nuevas de monitoreo."
    )
    add_table(document, ["Dato operativo", "Para que sirve"], [
        ["Identificador", "Cada respuesta devuelve X-Request-Id para localizar su registro."],
        ["Solicitudes", "Cuenta respuestas completadas, errores por clase, solicitudes activas y abortos."],
        ["Duraciones", "Mide media, maximo y distribucion acumulada del tiempo de respuesta."],
        ["Proceso", "Informa memoria residente y tiempo activo del monitor."],
    ], [1.55, 5.25])
    document.add_heading("Como se utiliza", level=2)
    add_bullets(document, [
        "Los registros JSON aparecen en la salida del servidor al atender solicitudes.",
        "El operador configura OPERATIONS_METRICS_TOKEN con 64 caracteres hexadecimales aleatorios y reinicia.",
        "Consulta GET /api/v1/operations/metrics con la clave en Authorization Bearer; fuera de localhost debe usar HTTPS.",
        "Sin clave configurada responde 404; con una clave incorrecta, 401. No se activo en la instalacion local del usuario.",
    ])
    document.add_heading("Privacidad y acceso", level=2)
    document.add_paragraph(
        "Los nuevos registros HTTP omiten URL, cuerpos, cookies, tokens, direcciones IP y datos "
        "de usuarios o empresas. El identificador lo genera el servidor. Las metricas son globales "
        "del proceso: solo el operador con la clave puede consultarlas; los roles de una empresa "
        "no dan acceso. La pagina publica de GitHub Pages no tiene este endpoint."
    )
    document.add_heading("Limites y siguiente paso", level=2)
    document.add_paragraph(
        "Los contadores reinician con el proceso y no incluyen historicos, trazas ni alertas. "
        "Compose tiene prevista rotacion de logs, aun pendiente de validacion en Docker. "
        "Faltan centralizacion, pruebas de carga, PostgreSQL, cola compartida, despliegue productivo "
        "y validacion social con cuentas reales. No se agregaron dependencias ni se modificaron "
        "secretos. Guia completa: docs/30-monitoreo-operativo.md."
    )


def main():
    for path in [MANUAL, WEEKLY]:
        document = Document(path)
        document.core_properties.author = "Social Audit Pro Team"
        document.core_properties.last_modified_by = "Social Audit Pro Team"
        if path == MANUAL:
            for table in document.tables:
                if table.rows and table.rows[0].cells[0].text == "Dato":
                    for row in table.rows:
                        if row.cells[0].text == "Version documentada":
                            row.cells[1].text = "v0.5.2"
                        if row.cells[0].text == "Fecha de corte":
                            row.cells[1].text = "26 de septiembre de 2026"
                        if row.cells[0].text == "Referencia de avance":
                            row.cells[1].text = "avance-013"
                    break
        else:
            for paragraph in document.paragraphs:
                if paragraph.text.startswith("Version documentada"):
                    paragraph.text = "Version documentada  v0.5.2  |  Fecha de corte  26 de septiembre de 2026"
                    break
        append_monitoring(document)
        document.save(path)
        print(path)


if __name__ == "__main__":
    main()
