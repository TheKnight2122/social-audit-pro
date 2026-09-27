from docx import Document
from update_v050_docs import MANUAL, WEEKLY

TITLE = "Primeras operaciones del backend en PostgreSQL"


def main():
    for path in [MANUAL, WEEKLY]:
        document = Document(path)
        if any(p.text == TITLE for p in document.paragraphs):
            continue
        document.core_properties.author = "Social Audit Pro Team"
        document.core_properties.last_modified_by = "Social Audit Pro Team"
        if path == MANUAL:
            for table in document.tables:
                if table.rows and table.rows[0].cells[0].text == "Dato":
                    for row in table.rows:
                        if row.cells[0].text == "Version documentada":
                            row.cells[1].text = "v0.5.7"
                        if row.cells[0].text == "Referencia de avance":
                            row.cells[1].text = "avance-018"
                    break
        else:
            for paragraph in document.paragraphs:
                if paragraph.text.startswith("Version documentada"):
                    paragraph.text = "Version documentada  v0.5.7  |  Fecha de corte  26 de septiembre de 2026"
                    break
        document.add_page_break()
        document.add_heading(TITLE, level=1)
        document.add_paragraph(
            "Avance 018 del 26 de septiembre de 2026. Se adapto la lectura de sesiones y el "
            "limite de intentos de acceso para trabajar con SQLite o PostgreSQL. El sistema "
            "completo todavia inicia con SQLite. Las pruebas PostgreSQL usan sesiones preparadas "
            "para validacion, no un registro o inicio de sesion completo del producto."
        )
        document.add_heading("Que mejora", level=2)
        document.add_paragraph(
            "El servidor sigue comprobando usuario, rol y empresa antes de aceptar la sesion. "
            "Ahora interpreta las fechas de vencimiento correctamente aunque tengan formato "
            "ISO o un desplazamiento horario. Conserva permisos y proteccion CSRF. Un fallo "
            "de base no concede acceso ni permite continuar al controlador protegido."
        )
        document.add_paragraph(
            "El limite agrupa conteo e insercion de intentos para evitar que solicitudes "
            "simultaneas superen el maximo. PostgreSQL utiliza un bloqueo compartido por clave. "
            "Tambien se preparo un grupo acotado de conexiones, con TLS validado, UTC, tiempos "
            "maximos y transacciones que liberan o descartan la conexion segun el resultado."
        )
        document.add_heading("Pruebas y continuidad", level=2)
        document.add_paragraph(
            "Pasaron 96 pruebas automatizadas y la construccion publica. La carga local "
            "ejecuto 200 solicitudes sin errores. La prueba del servidor PostgreSQL en GitHub "
            "usa dos grupos de conexiones y 30 intentos concurrentes para comprobar el limite. "
            "La version local conserva sus funciones y la publica sigue siendo una demo."
        )
        document.add_heading("Trabajo que sigue pendiente", level=2)
        document.add_paragraph(
            "Falta adaptar a PostgreSQL las demas operaciones de autenticacion, organizaciones, "
            "analitica, reportes e integraciones y los trabajadores. Esas funciones siguen "
            "disponibles con SQLite. Despues se podra seleccionar el motor y desplegar el piloto. "
            "No se trasladaron datos reales ni se contrataron servicios. Correo y redes "
            "oficiales siguen pendientes de autorizaciones. Guia: docs/35-acceso-asincrono-postgresql.md."
        )
        document.save(path)
        print(path)


if __name__ == "__main__":
    main()
