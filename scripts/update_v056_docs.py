from docx import Document
from update_v050_docs import MANUAL, WEEKLY

TITLE = "Preparacion del traslado a PostgreSQL"


def main():
    for path in [MANUAL, WEEKLY]:
        document = Document(path)
        if any(p.text == TITLE for p in document.paragraphs):
            continue
        document.core_properties.author = "Social Audit Pro Team"
        document.core_properties.last_modified_by = "Social Audit Pro Team"
        if path == MANUAL:
            for table in document.tables:
                if table.rows and table.rows[0].cells[0].text == "Dato":
                    for row in table.rows:
                        if row.cells[0].text == "Version documentada":
                            row.cells[1].text = "v0.5.6"
                        if row.cells[0].text == "Referencia de avance":
                            row.cells[1].text = "avance-017"
                    break
        else:
            for paragraph in document.paragraphs:
                if paragraph.text.startswith("Version documentada"):
                    paragraph.text = "Version documentada  v0.5.6  |  Fecha de corte  26 de septiembre de 2026"
                    break
        document.add_page_break()
        document.add_heading(TITLE, level=1)
        document.add_paragraph(
            "Avance 017 del 26 de septiembre de 2026. Se agrego una herramienta para trasladar "
            "datos de SQLite a una base PostgreSQL nueva. El traslado se verifica antes de "
            "confirmarse. La aplicacion local sigue usando SQLite: todavia no se ha adaptado "
            "el backend completo ni se ha desplegado en internet."
        )
        document.add_heading("Que conserva y que comprueba", level=2)
        document.add_paragraph(
            "Conserva empresas, usuarios, permisos, cuentas, publicaciones, metricas, reportes "
            "y secretos cifrados. Comprueba referencias y contenido, rechaza destinos ocupados "
            "y cancela la transaccion si falla una comprobacion. Trabaja sobre una instantanea "
            "en memoria, sin modificar la base original. Las claves de cifrado deben custodiarse aparte."
        )
        document.add_heading("Medidas antes de abrir el nuevo entorno", level=2)
        document.add_paragraph(
            "En el destino se invalidan sesiones y enlaces temporales, se pausan sincronizaciones "
            "y correos pendientes, y se reinician los codigos de recuperacion MFA. El segundo "
            "factor ya habilitado se conserva. El administrador debera verificar acceso, renovar "
            "codigos y revisar tareas antes de habilitar usuarios. No se activa ninguna API de pago."
        )
        document.add_heading("Evidencia y pendientes", level=2)
        document.add_paragraph(
            "Pasaron 84 pruebas automatizadas, incluidas ocho nuevas con PostgreSQL embebido, "
            "SQLite y la herramienta de consola. La construccion publica paso y la carga local "
            "de 200 solicitudes no tuvo errores. El workflow incorpora una prueba adicional "
            "contra PostgreSQL 16 servidor con datos desechables y dos conexiones."
        )
        document.add_paragraph(
            "Falta adaptar las consultas y transacciones del backend a PostgreSQL, desplegarlo "
            "y activar correo y redes con autorizaciones reales. No se contrataron servicios "
            "ni se trasladaron datos reales. La pagina publica sigue siendo una demo. "
            "Guia operativa: docs/34-traslado-postgresql.md."
        )
        document.save(path)
        print(path)


if __name__ == "__main__":
    main()
