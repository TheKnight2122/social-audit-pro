from pathlib import Path
from docx import Document
from docx.shared import Inches
from update_v050_docs import MANUAL, WEEKLY

TITLE = "Acceso de clientes y actualizacion del dashboard"


def main():
    for path in [MANUAL, WEEKLY]:
        document = Document(path)
        if any(p.text == TITLE for p in document.paragraphs):
            continue
        document.core_properties.author = "Social Audit Pro Team"
        document.core_properties.last_modified_by = "Social Audit Pro Team"
        if path == MANUAL:
            for table in document.tables:
                if table.rows and table.rows[0].cells[0].text == "Dato":
                    for row in table.rows:
                        if row.cells[0].text == "Version documentada":
                            row.cells[1].text = "v0.5.4"
                        if row.cells[0].text == "Referencia de avance":
                            row.cells[1].text = "avance-015"
                    break
        else:
            for paragraph in document.paragraphs:
                if paragraph.text.startswith("Version documentada"):
                    paragraph.text = "Version documentada  v0.5.4  |  Fecha de corte  26 de septiembre de 2026"
                    break
        document.add_page_break()
        document.add_heading(TITLE, level=1)
        document.add_paragraph(
            "Avance 015 del 26 de septiembre de 2026. Se prepararon dos entradas de acceso y "
            "un resumen de resultados para clientes. El servidor actualiza las conexiones al "
            "ingresar al dashboard, con un intervalo minimo de quince minutos por conexion. "
            "Este avance no despliega el backend online ni activa redes sin sus credenciales."
        )
        document.add_heading("Uso y permisos", level=2)
        document.add_paragraph(
            "Los enlaces #/cuenta/admin y #/cuenta/client comparten la autenticacion existente. "
            "El servidor decide el rol real; elegir Administradores no concede privilegios. "
            "El administrador crea usuarios en la organizacion correspondiente desde Configuracion. "
            "Se propone una organizacion por empresa cliente, pendiente de confirmar si tambien "
            "se necesitan restricciones por usuario dentro de la empresa."
        )
        document.add_paragraph(
            "El cliente ve comunidad, publicaciones, visualizaciones, likes y una tabla de impacto. "
            "Las metricas ausentes dicen No disponible. Sin cuentas conectadas no se muestran "
            "cifras demostrativas. Cuenta conserva las opciones de seguridad. La simplificacion "
            "visual no elimina los permisos analiticos de lectura de su organizacion por API."
        )
        document.add_heading("Actualizacion y validacion", level=2)
        document.add_paragraph(
            "El refresco usa sesion, CSRF, organizacion activa y el bloqueo compartido. "
            "Se limita cada consulta a veinte segundos; ante fallo se conservan los datos anteriores. "
            "Los horarios existentes no se desactivan. Se aprobaron 66 pruebas automatizadas y "
            "se revisaron ambos accesos y el dashboard cliente en escritorio y movil con datos desechables."
        )
        document.add_heading("Decisiones y pendientes", level=2)
        document.add_paragraph(
            "Se solicito alojamiento gratuito, prueba con datos reales y posterior uso diario. "
            "La gratuidad no garantiza continuidad ni APIs gratuitas para todas las redes. "
            "Existe correo empresarial, pero falta confirmar proveedor y configurar el envio. "
            "Siguen pendientes PostgreSQL, alojamiento, permisos sociales y validacion productiva. "
            "No se contrataron servicios. GitHub Pages sigue siendo una demostracion estatica."
        )
        document.add_page_break()
        document.add_heading("Vista de resultados para clientes", level=1)
        document.add_paragraph(
            "Captura de verificacion con cuenta y proveedor simulados. El rotulo Datos oficiales "
            "corresponde al estado de interfaz probado, no a una conexion real activada. "
            "La metrica ausente se identifica sin inventar valores."
        )
        document.add_picture(str(Path(__file__).resolve().parents[1] / "docs/assets/portal/cliente-desktop.png"), width=Inches(6.3))
        document.add_paragraph(
            "La revision movil se conserva en docs/assets/portal/cliente-mobile.png. "
            "Los dos accesos se documentan con las capturas de la misma carpeta. "
            "La guia tecnica completa esta en docs/32-portales-y-actualizacion-dashboard.md."
        )
        document.save(path)
        print(path)


if __name__ == "__main__":
    main()
